"""Export the measured Markdown report as a readable Word document."""
from pathlib import Path
import re
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
TITLE = '高效率BERT古文断句标点微调与批量长文推理实验'


def inline(paragraph, tokens):
    strong = False
    link = None
    for token in tokens or []:
        if token.type == 'strong_open':
            strong = True
        elif token.type == 'strong_close':
            strong = False
        elif token.type in ['text', 'code_inline', 'softbreak', 'hardbreak']:
            if link is not None:
                element = OxmlElement('w:hyperlink')
                element.set(qn('r:id'), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
                run_element = OxmlElement('w:r')
                properties = OxmlElement('w:rPr')
                color = OxmlElement('w:color')
                color.set(qn('w:val'), '285B90')
                properties.append(color)
                run_element.append(properties)
                content = OxmlElement('w:t')
                content.text = token.content
                run_element.append(content)
                element.append(run_element)
                paragraph._p.append(element)
                continue
            run = paragraph.add_run('\n' if token.type in ['softbreak', 'hardbreak'] else token.content)
            run.bold = strong
            if token.type == 'code_inline':
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
        elif token.type == 'link_open':
            link = token.attrGet('href')
        elif token.type == 'link_close':
            link = None


def setup(document):
    sec = document.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    # Remove borders inherited from the runtime's base Word template.
    for borders in list(document.styles.element.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    document.styles['Title'].font.size = Pt(21)
    for name, size in [('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 11.5)]:
        style = document.styles[name]
        style.font.size, style.font.bold = Pt(size), True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    document.styles['Caption'].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)


def format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    n = len(table.columns)
    widths = [16.8 / n] * n
    if n == 7:
        widths = [1.45, 2.25, 1.3, 2.5, 3, 3, 3.3]
        widths = [w * 16.8 / sum(widths) for w in widths]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        item = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            item.set(qn('w:' + key), value)
        borders.append(item)
    table._tbl.tblPr.append(borders)
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DBE7F1' if i == 0 else 'FFFFFF' if i % 2 else 'F4F7FA')
            cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if i == 0:
                    p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    run.bold = i == 0


def main():
    markdown = (ROOT / '作业报告.md').read_text(encoding='utf-8')
    tokens = MarkdownIt('commonmark').enable('table').parse(markdown)
    doc = Document()
    setup(doc)
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type == 'heading_open':
            level = int(token.tag[1:])
            paragraph = doc.add_paragraph(style='Title' if level == 1 else f'Heading {min(level-1,3)}')
            inline(paragraph, tokens[i + 1].children)
            if level == 1:
                paragraph.clear()
                paragraph.add_run('高效率 BERT 古文断句标点\n微调与批量长文推理实验')
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 3
            continue
        if token.type == 'paragraph_open':
            children = tokens[i + 1].children or []
            images = [t for t in children if t.type == 'image']
            if images:
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                for picture in images:
                    source = picture.attrGet('src')
                    if source.startswith('https://raw.githubusercontent.com/Can1017/LLM-HW/main/hw3/'):
                        source = unquote(urlparse(source).path.split('/main/hw3/', 1)[1])
                    paragraph.add_run().add_picture(str(ROOT / unquote(source)), width=Cm(16.5))
            else:
                paragraph = doc.add_paragraph()
                inline(paragraph, children)
                if re.match(r'^图\s*\d', paragraph.text):
                    paragraph.style = 'Caption'
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 3
            continue
        if token.type == 'fence':
            for line in token.content.rstrip().splitlines():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                r = p.add_run(line)
                r.font.name, r.font.size = 'Consolas', Pt(8)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif token.type == 'table_open':
            rows, row, column = [], None, None
            i += 1
            while tokens[i].type != 'table_close':
                t = tokens[i]
                if t.type == 'tr_open':
                    row = []
                    rows.append(row)
                elif t.type in ['th_open', 'td_open']:
                    row.append(tokens[i + 1])
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for row_cells, contents in zip(table.rows, rows):
                for cell, content in zip(row_cells.cells, contents):
                    inline(cell.paragraphs[0], content.children)
            format_table(table)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif token.type == 'inline':
            paragraph = doc.add_paragraph()
            inline(paragraph, token.children)
        i += 1
    destination = ROOT / (TITLE + '.docx')
    doc.save(destination)
    print(destination)


if __name__ == '__main__':
    main()
